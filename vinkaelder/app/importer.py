"""Import fra den gamle Excel-vindatabase (.xlsm/.xlsx).

Kan køres igen og igen:
- Vine matches på stregkode (ellers beskrivelse + årgang) og opdateres.
  Felter, der kun findes i appen (vintype, placering), bevares.
- Alle bevægelser med kilde 'excel' erstattes af arkets Lager-fane.
  Bevægelser oprettet i appen bevares, medmindre drop_app=True.
"""
import io
import re
from datetime import datetime, date

import openpyxl

from db import now
from estimate import norm

HEADER_MAP = {
    "beskrivelse": "description", "årgang": "vintage", "lager": "_excel_stock", "land": "country",
    "producent": "producer", "distrikt": "region", "drink age from": "drink_from_age",
    "drink age to": "drink_to_age", "druer": "grapes", "anskaffelsespris": "purchase_price",
    "listepris": "list_price", "barcode": "barcode", "vinhandel": "shop", "note": "note", "link": "link",
}
# Kolonne U-AA (Drink from/to (years after vintage), drinking year from/to, Drinking advice,
# rationale, advice2) importeres bevidst ikke.
EXCEL_OWNED = ["barcode", "description", "vintage", "country", "producer", "region", "grapes", "drink_from_age",
               "drink_to_age", "purchase_price", "list_price", "shop", "note", "link"]


def clean_barcode(v):
    if v is None:
        return None
    if isinstance(v, float):
        v = int(v)
    s = str(v).strip()
    if s in ("", "0", "NA", "None"):
        return None
    return s


def clean_text(v):
    if v is None:
        return None
    s = str(v).replace("_x000D_", "").replace("\r", "").strip()
    return s or None


def clean_grapes(v):
    s = clean_text(v)
    if not s:
        return None
    parts = [p.strip() for p in re.split(r"[\n,;]+", s) if p.strip()]
    return ", ".join(parts)


def to_int(v):
    try:
        return int(round(float(v)))
    except (TypeError, ValueError):
        return None


def to_float(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def read_workbook(data):
    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    if "Vin" not in wb.sheetnames or "Lager" not in wb.sheetnames:
        raise ValueError("Filen skal have fanerne 'Vin' og 'Lager'.")
    rows = list(wb["Vin"].iter_rows(values_only=True))
    header = [str(h).strip().lower() if h is not None else "" for h in rows[0]]
    idx = {}
    for i, h in enumerate(header):
        if h in HEADER_MAP and HEADER_MAP[h] not in idx:
            idx[HEADER_MAP[h]] = i
    if "description" not in idx:
        raise ValueError("Kunne ikke finde kolonnen 'Beskrivelse' i fanen Vin.")
    wines = []
    for r in rows[1:]:
        g = lambda k: r[idx[k]] if k in idx and idx[k] < len(r) else None  # noqa: E731
        desc = clean_text(g("description"))
        if not desc:
            continue
        w = {
            "description": desc, "vintage": to_int(g("vintage")), "country": clean_text(g("country")),
            "producer": clean_text(g("producer")), "region": clean_text(g("region")),
            "grapes": clean_grapes(g("grapes")), "drink_from_age": to_int(g("drink_from_age")),
            "drink_to_age": to_int(g("drink_to_age")),
            "purchase_price": to_float(g("purchase_price")), "list_price": to_float(g("list_price")),
            "barcode": clean_barcode(g("barcode")), "shop": clean_text(g("shop")), "note": clean_text(g("note")),
            "link": clean_text(g("link")) if str(g("link") or "").startswith("http") else None,
            "_excel_stock": to_int(g("_excel_stock")),
        }
        wines.append(w)

    moves = []
    for r in list(wb["Lager"].iter_rows(values_only=True))[1:]:
        if not r or len(r) < 4:
            continue
        d, bc, qty, act = r[0], r[1], r[2], r[3]
        bc = clean_barcode(bc)
        if bc is None or not isinstance(d, (datetime, date)):
            continue
        qty = to_int(qty) or 0
        act = str(act or "").strip()
        m = re.match(r"status\s*\((-?\d+)", act, re.I)
        if m:
            kind, count_value = "count", int(m.group(1))
        elif act.lower() == "in":
            kind, count_value = "in", None
        elif act.lower() == "out":
            kind, count_value = "out", None
        else:
            kind, count_value = ("in" if qty >= 0 else "out"), None
        note = None if kind != "count" and act.lower() in ("in", "out") else (act if kind != "count" else None)
        moves.append({"date": d.strftime("%Y-%m-%d"), "barcode": bc, "qty": qty, "kind": kind,
                      "count_value": count_value, "note": note})
    return wines, moves


def assign_keys(wines):
    from collections import Counter
    bc_count = Counter(w["barcode"] for w in wines if w["barcode"])
    for w in wines:
        if w["barcode"] and bc_count[w["barcode"]] == 1:
            w["import_key"] = "bc:" + w["barcode"]
        elif w["barcode"]:
            w["import_key"] = f"bc:{w['barcode']}:{w['vintage']}:{norm(w['description'])}"
        else:
            w["import_key"] = f"d:{norm(w['description'])}:{w['vintage']}"
    # sidste udvej: gør dubletter unikke
    seen = Counter()
    for w in wines:
        seen[w["import_key"]] += 1
        if seen[w["import_key"]] > 1:
            w["import_key"] += f"#{seen[w['import_key']]}"
    return bc_count


def run_import(con, data, dry_run=False, drop_app=False):
    wines, moves = read_workbook(data)
    bc_count = assign_keys(wines)
    summary = {"wines_in_file": len(wines), "wines_new": 0, "wines_updated": 0, "wines_unchanged": 0,
               "moves_in_file": len(moves), "moves_imported": 0, "orphan_barcodes": [], "orphan_moves": 0,
               "app_moves_kept": 0, "app_moves_deleted": 0, "stock_mismatch": [], "wines_not_in_file": [],
               "dry_run": dry_run}
    ts = now()
    con.execute("BEGIN")
    try:
        key_to_id = {}
        for w in wines:
            row = con.execute("SELECT * FROM wine WHERE import_key=?", (w["import_key"],)).fetchone()
            if row is None and w["barcode"] and bc_count[w["barcode"]] == 1:
                # vin oprettet i appen med samme stregkode -> knyt sammen
                row = con.execute("SELECT * FROM wine WHERE barcode=? AND import_key IS NULL", (w["barcode"],)).fetchone()
            if row is None:
                cols = EXCEL_OWNED + ["import_key", "created_at", "updated_at"]
                vals = [w[c] for c in EXCEL_OWNED] + [w["import_key"], ts, ts]
                cur = con.execute(f"INSERT INTO wine ({','.join(cols)}) VALUES ({','.join('?' * len(cols))})", vals)
                key_to_id[w["import_key"]] = cur.lastrowid
                summary["wines_new"] += 1
            else:
                changed = [c for c in EXCEL_OWNED if row[c] != w[c]]
                if changed or row["import_key"] != w["import_key"]:
                    sets = ",".join(f"{c}=?" for c in EXCEL_OWNED)
                    con.execute(f"UPDATE wine SET {sets}, import_key=?, updated_at=? WHERE id=?",
                                [w[c] for c in EXCEL_OWNED] + [w["import_key"], ts, row["id"]])
                    summary["wines_updated"] += 1
                else:
                    summary["wines_unchanged"] += 1
                key_to_id[w["import_key"]] = row["id"]

        keys = {w["import_key"] for w in wines}
        for r in con.execute("SELECT id, description, import_key FROM wine WHERE import_key IS NOT NULL"):
            if r["import_key"] not in keys:
                summary["wines_not_in_file"].append(r["description"])

        # bevægelser
        con.execute("DELETE FROM movement WHERE source='excel'")
        if drop_app:
            summary["app_moves_deleted"] = con.execute("DELETE FROM movement WHERE source='app'").rowcount
        summary["app_moves_kept"] = con.execute("SELECT COUNT(*) FROM movement WHERE source='app'").fetchone()[0]
        bc_to_id = {}
        for w in wines:
            if w["barcode"] and bc_count[w["barcode"]] == 1:
                bc_to_id[w["barcode"]] = key_to_id[w["import_key"]]
        # vine oprettet i appen (uden import_key) kan også modtage bevægelser
        for r in con.execute("SELECT id, barcode FROM wine WHERE import_key IS NULL AND barcode IS NOT NULL"):
            bc_to_id.setdefault(r["barcode"], r["id"])
        orphans = set()
        for m in moves:
            wid = bc_to_id.get(m["barcode"])
            if wid is None:
                orphans.add(m["barcode"])
                summary["orphan_moves"] += 1
            con.execute("INSERT INTO movement (wine_id, barcode, date, qty, kind, count_value, note, source, created_at)"
                        " VALUES (?,?,?,?,?,?,?,'excel',?)",
                        (wid, m["barcode"], m["date"], m["qty"], m["kind"], m["count_value"], m["note"], ts))
            summary["moves_imported"] += 1
        summary["orphan_barcodes"] = sorted(orphans)

        # kontrol: lagertal i appen vs. Excel-kolonnen 'Lager'
        for w in wines:
            if w["_excel_stock"] is None:
                continue
            wid = key_to_id[w["import_key"]]
            s = con.execute("SELECT COALESCE(SUM(qty),0) FROM movement WHERE wine_id=?", (wid,)).fetchone()[0]
            if s != w["_excel_stock"]:
                summary["stock_mismatch"].append({"description": w["description"], "excel": w["_excel_stock"], "app": s})

        if dry_run:
            con.execute("ROLLBACK")
        else:
            con.execute("INSERT OR REPLACE INTO meta(key,value) VALUES('last_import',?)", (ts,))
            con.execute("COMMIT")
    except Exception:
        con.execute("ROLLBACK")
        raise
    return summary
